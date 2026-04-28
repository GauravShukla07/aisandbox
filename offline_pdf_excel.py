from __future__ import annotations

import argparse
import logging
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

import cv2
import numpy as np
import pypdfium2 as pdfium
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from rapidocr_onnxruntime import RapidOCR


LOGGER = logging.getLogger("offline_pdf_excel")


@dataclass(frozen=True)
class TableRegion:
    bbox: tuple[int, int, int, int]
    vertical_lines: list[int]
    horizontal_lines: list[int]


@dataclass(frozen=True)
class ExtractedTable:
    page_index: int
    table_index: int
    rows: list[list[str]]


def render_pdf_pages(pdf_path: Path, dpi: int) -> list[np.ndarray]:
    document = pdfium.PdfDocument(str(pdf_path))
    scale = dpi / 72.0
    pages: list[np.ndarray] = []

    for page_index in range(len(document)):
        page = document[page_index]
        bitmap = page.render(scale=scale)
        image = bitmap.to_numpy().copy()
        pages.append(normalize_bgr_image(image))

    return pages


def normalize_bgr_image(image: np.ndarray) -> np.ndarray:
    if image.ndim == 2:
        return cv2.cvtColor(image, cv2.COLOR_GRAY2BGR)
    if image.shape[2] == 4:
        return cv2.cvtColor(image, cv2.COLOR_BGRA2BGR)
    return image


def binarize_image(image: np.ndarray) -> np.ndarray:
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    blurred = cv2.GaussianBlur(gray, (3, 3), 0)
    return cv2.adaptiveThreshold(
        255 - blurred,
        255,
        cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
        cv2.THRESH_BINARY,
        31,
        11,
    )


def detect_grid_masks(binary: np.ndarray) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    height, width = binary.shape
    horizontal_kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT, (max(20, width // 30), 1)
    )
    vertical_kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT, (1, max(20, height // 30))
    )

    horizontal = cv2.morphologyEx(binary, cv2.MORPH_OPEN, horizontal_kernel)
    vertical = cv2.morphologyEx(binary, cv2.MORPH_OPEN, vertical_kernel)
    horizontal = cv2.dilate(
        horizontal, cv2.getStructuringElement(cv2.MORPH_RECT, (3, 1)), iterations=1
    )
    vertical = cv2.dilate(
        vertical, cv2.getStructuringElement(cv2.MORPH_RECT, (1, 3)), iterations=1
    )
    grid = cv2.add(horizontal, vertical)
    return horizontal, vertical, grid


def detect_table_regions(
    grid_mask: np.ndarray,
    min_area_ratio: float,
    min_width: int = 80,
    min_height: int = 80,
) -> list[tuple[int, int, int, int]]:
    image_height, image_width = grid_mask.shape
    image_area = image_height * image_width

    grown = cv2.dilate(
        grid_mask,
        cv2.getStructuringElement(cv2.MORPH_RECT, (15, 15)),
        iterations=2,
    )
    contours, _ = cv2.findContours(grown, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)

    boxes: list[tuple[int, int, int, int]] = []
    for contour in contours:
        x, y, width, height = cv2.boundingRect(contour)
        if width < min_width or height < min_height:
            continue
        if width * height < image_area * min_area_ratio:
            continue
        boxes.append((x, y, width, height))

    if not boxes:
        return [(0, 0, image_width, image_height)]

    boxes = merge_boxes(boxes, gap=12)
    boxes.sort(key=lambda box: (box[1], box[0]))
    return boxes


def merge_boxes(
    boxes: Sequence[tuple[int, int, int, int]], gap: int
) -> list[tuple[int, int, int, int]]:
    if not boxes:
        return []

    remaining = [tuple(box) for box in boxes]
    merged = True
    while merged:
        merged = False
        next_round: list[tuple[int, int, int, int]] = []
        while remaining:
            current = remaining.pop()
            cx, cy, cw, ch = current
            changed = True
            while changed:
                changed = False
                current_left = cx - gap
                current_top = cy - gap
                current_right = cx + cw + gap
                current_bottom = cy + ch + gap

                for index in range(len(remaining) - 1, -1, -1):
                    ox, oy, ow, oh = remaining[index]
                    other_left = ox
                    other_top = oy
                    other_right = ox + ow
                    other_bottom = oy + oh

                    if (
                        current_left <= other_right
                        and current_right >= other_left
                        and current_top <= other_bottom
                        and current_bottom >= other_top
                    ):
                        nx1 = min(cx, ox)
                        ny1 = min(cy, oy)
                        nx2 = max(cx + cw, ox + ow)
                        ny2 = max(cy + ch, oy + oh)
                        cx, cy, cw, ch = nx1, ny1, nx2 - nx1, ny2 - ny1
                        remaining.pop(index)
                        changed = True
                        merged = True
            next_round.append((cx, cy, cw, ch))
        remaining = next_round
    return remaining


def cluster_positions(indices: Iterable[int]) -> list[int]:
    values = [int(index) for index in indices]
    if not values:
        return []

    groups: list[list[int]] = [[values[0]]]
    for value in values[1:]:
        if value - groups[-1][-1] <= 1:
            groups[-1].append(value)
            continue
        groups.append([value])

    return [int(round(sum(group) / len(group))) for group in groups]


def merge_close_positions(positions: Sequence[int], min_gap: int) -> list[int]:
    if not positions:
        return []

    merged = [int(positions[0])]
    for value in positions[1:]:
        if value - merged[-1] <= min_gap:
            merged[-1] = int(round((merged[-1] + value) / 2))
        else:
            merged.append(int(value))
    return merged


def extract_line_positions(
    line_mask: np.ndarray, axis: str, coverage_ratio: float, min_gap: int
) -> list[int]:
    if axis == "vertical":
        projection = np.count_nonzero(line_mask, axis=0)
        threshold = max(5, int(line_mask.shape[0] * coverage_ratio))
    elif axis == "horizontal":
        projection = np.count_nonzero(line_mask, axis=1)
        threshold = max(5, int(line_mask.shape[1] * coverage_ratio))
    else:
        raise ValueError(f"Unsupported axis: {axis}")

    indices = np.flatnonzero(projection >= threshold)
    return merge_close_positions(cluster_positions(indices.tolist()), min_gap=min_gap)


def infer_lines_from_cell_contours(
    grid_mask: np.ndarray,
    min_cell_width: int,
    min_cell_height: int,
    min_gap: int,
) -> tuple[list[int], list[int]]:
    contours, _ = cv2.findContours(grid_mask, cv2.RETR_TREE, cv2.CHAIN_APPROX_SIMPLE)
    mask_height, mask_width = grid_mask.shape
    min_area = max(min_cell_width * min_cell_height * 4, 4000)

    x_positions: list[int] = []
    y_positions: list[int] = []
    for contour in contours:
        x, y, width, height = cv2.boundingRect(contour)
        if width < min_cell_width or height < min_cell_height:
            continue
        if width * height < min_area:
            continue
        if width >= mask_width * 0.98 and height >= mask_height * 0.98:
            continue

        x_positions.extend((x, x + width))
        y_positions.extend((y, y + height))

    vertical_lines = merge_close_positions(sorted(x_positions), min_gap=min_gap)
    horizontal_lines = merge_close_positions(sorted(y_positions), min_gap=min_gap)
    return vertical_lines, horizontal_lines


def prepare_region(
    image: np.ndarray,
    horizontal_mask: np.ndarray,
    vertical_mask: np.ndarray,
    bbox: tuple[int, int, int, int],
    coverage_ratio: float,
    min_cell_width: int,
    min_cell_height: int,
) -> TableRegion | None:
    x, y, width, height = bbox
    region_horizontal = horizontal_mask[y : y + height, x : x + width]
    region_vertical = vertical_mask[y : y + height, x : x + width]
    region_grid = cv2.add(region_horizontal, region_vertical)

    min_gap = max(4, min(width, height) // 200)
    vertical_lines = extract_line_positions(
        region_vertical, axis="vertical", coverage_ratio=coverage_ratio, min_gap=min_gap
    )
    horizontal_lines = extract_line_positions(
        region_horizontal,
        axis="horizontal",
        coverage_ratio=coverage_ratio,
        min_gap=min_gap,
    )

    if len(vertical_lines) < 2 or len(horizontal_lines) < 2:
        fallback_vertical, fallback_horizontal = infer_lines_from_cell_contours(
            region_grid,
            min_cell_width=min_cell_width,
            min_cell_height=min_cell_height,
            min_gap=max(8, min_gap * 2),
        )
        if len(fallback_vertical) > len(vertical_lines):
            vertical_lines = fallback_vertical
        if len(fallback_horizontal) > len(horizontal_lines):
            horizontal_lines = fallback_horizontal

    if len(vertical_lines) < 2 or len(horizontal_lines) < 2:
        LOGGER.debug("Skipping region without stable grid: %s", bbox)
        return None

    return TableRegion(
        bbox=bbox,
        vertical_lines=vertical_lines,
        horizontal_lines=horizontal_lines,
    )


def preprocess_cell_for_ocr(cell_image: np.ndarray) -> np.ndarray:
    gray = cv2.cvtColor(cell_image, cv2.COLOR_BGR2GRAY)
    bordered = cv2.copyMakeBorder(
        gray, 6, 6, 6, 6, cv2.BORDER_CONSTANT, value=255
    )
    scaled = cv2.resize(
        bordered, None, fx=2.0, fy=2.0, interpolation=cv2.INTER_CUBIC
    )
    thresholded = cv2.threshold(
        scaled, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU
    )[1]
    return cv2.cvtColor(thresholded, cv2.COLOR_GRAY2BGR)


def parse_ocr_text(
    ocr_result: list[list[object]] | None, score_threshold: float
) -> str:
    if not ocr_result:
        return ""

    texts: list[str] = []
    for entry in ocr_result:
        if len(entry) >= 3:
            _, text, score = entry[:3]
        elif len(entry) == 2:
            text, score = entry
        else:
            continue

        if float(score) < score_threshold:
            continue
        cleaned = str(text).strip()
        if cleaned:
            texts.append(cleaned)

    return " ".join(texts).strip()


def ocr_cell(
    engine: RapidOCR, cell_image: np.ndarray, score_threshold: float
) -> str:
    prepared = preprocess_cell_for_ocr(cell_image)

    rec_only_result, _ = engine(
        prepared, use_det=False, use_cls=False, use_rec=True
    )
    text = parse_ocr_text(rec_only_result, score_threshold=score_threshold)
    if text:
        return text

    detected_result, _ = engine(prepared)
    return parse_ocr_text(detected_result, score_threshold=score_threshold)


def extract_rows_from_region(
    image: np.ndarray,
    region: TableRegion,
    min_cell_width: int,
    min_cell_height: int,
    cell_padding: int,
    engine: RapidOCR,
    ocr_score_threshold: float,
) -> list[list[str]]:
    x, y, width, height = region.bbox
    table_image = image[y : y + height, x : x + width]
    rows: list[list[str]] = []

    for top, bottom in zip(region.horizontal_lines, region.horizontal_lines[1:]):
        row_values: list[str] = []
        for left, right in zip(region.vertical_lines, region.vertical_lines[1:]):
            if right - left < min_cell_width or bottom - top < min_cell_height:
                continue

            crop_left = min(max(left + cell_padding, 0), width)
            crop_top = min(max(top + cell_padding, 0), height)
            crop_right = min(max(right - cell_padding, crop_left + 1), width)
            crop_bottom = min(max(bottom - cell_padding, crop_top + 1), height)
            cell = table_image[crop_top:crop_bottom, crop_left:crop_right]
            row_values.append(ocr_cell(engine, cell, score_threshold=ocr_score_threshold))

        if any(value.strip() for value in row_values):
            rows.append(row_values)

    return trim_empty_edges(rows)


def trim_empty_edges(rows: list[list[str]]) -> list[list[str]]:
    if not rows:
        return []

    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]

    leading = 0
    trailing = width
    while leading < trailing and all(not row[leading].strip() for row in normalized):
        leading += 1
    while trailing > leading and all(
        not row[trailing - 1].strip() for row in normalized
    ):
        trailing -= 1

    trimmed = [row[leading:trailing] for row in normalized]
    if not trimmed:
        return []

    width = max(len(row) for row in trimmed)
    normalized_trimmed = [row + [""] * (width - len(row)) for row in trimmed]
    keep_indices = [
        index
        for index in range(width)
        if any(row[index].strip() for row in normalized_trimmed)
    ]
    if not keep_indices:
        return []
    return [[row[index] for index in keep_indices] for row in normalized_trimmed]


def extract_tables(
    pdf_path: Path,
    dpi: int,
    region_area_ratio: float,
    line_coverage_ratio: float,
    min_cell_width: int,
    min_cell_height: int,
    cell_padding: int,
    ocr_score_threshold: float,
) -> list[ExtractedTable]:
    pages = render_pdf_pages(pdf_path, dpi=dpi)
    engine = RapidOCR()
    extracted: list[ExtractedTable] = []

    for page_index, image in enumerate(pages, start=1):
        binary = binarize_image(image)
        horizontal_mask, vertical_mask, grid_mask = detect_grid_masks(binary)
        regions = detect_table_regions(
            grid_mask, min_area_ratio=region_area_ratio, min_width=80, min_height=80
        )

        page_tables = 0
        for bbox in regions:
            region = prepare_region(
                image=image,
                horizontal_mask=horizontal_mask,
                vertical_mask=vertical_mask,
                bbox=bbox,
                coverage_ratio=line_coverage_ratio,
                min_cell_width=min_cell_width,
                min_cell_height=min_cell_height,
            )
            if region is None:
                continue

            rows = extract_rows_from_region(
                image=image,
                region=region,
                min_cell_width=min_cell_width,
                min_cell_height=min_cell_height,
                cell_padding=cell_padding,
                engine=engine,
                ocr_score_threshold=ocr_score_threshold,
            )
            if not rows:
                continue

            page_tables += 1
            extracted.append(
                ExtractedTable(
                    page_index=page_index, table_index=page_tables, rows=rows
                )
            )

    return extracted


def sheet_name_for(table: ExtractedTable) -> str:
    return truncate_sheet_title(f"page-{table.page_index}-table-{table.table_index}")


def truncate_sheet_title(value: str) -> str:
    cleaned = re.sub(r"[:\\\\/?*\\[\\]]", "-", value)
    return cleaned[:31]


def maybe_number(value: str, infer_numbers: bool) -> str | int | float:
    if not infer_numbers:
        return value

    stripped = value.strip()
    if not stripped:
        return value
    if re.fullmatch(r"-?\d+", stripped):
        if stripped.startswith("0") and len(stripped) > 1:
            return value
        if stripped.startswith("-0") and len(stripped) > 2:
            return value
        return int(stripped)
    if re.fullmatch(r"-?\d+\.\d+", stripped):
        if stripped.startswith("0") and len(stripped.split(".", maxsplit=1)[0]) > 1:
            return value
        return float(stripped)
    return value


def autosize_columns(worksheet) -> None:
    for column in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            if cell.value is None:
                continue
            max_length = max(max_length, len(str(cell.value)))
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 40)


def write_workbook(
    tables: Sequence[ExtractedTable], output_path: Path, infer_numbers: bool
) -> None:
    workbook = Workbook()
    first_sheet = True

    if not tables:
        worksheet = workbook.active
        worksheet.title = "page-1-table-1"
        worksheet["A1"] = "No table grid was detected."
        workbook.save(output_path)
        return

    for table in tables:
        worksheet = workbook.active if first_sheet else workbook.create_sheet()
        first_sheet = False
        worksheet.title = sheet_name_for(table)
        for row_index, row_values in enumerate(table.rows, start=1):
            for column_index, value in enumerate(row_values, start=1):
                worksheet.cell(
                    row=row_index,
                    column=column_index,
                    value=maybe_number(value, infer_numbers=infer_numbers),
                )
        autosize_columns(worksheet)

    workbook.save(output_path)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Extract Excel-like table data from scanned PDF pages entirely offline."
        )
    )
    parser.add_argument("input_pdf", type=Path, help="Path to the source PDF")
    parser.add_argument("output_xlsx", type=Path, help="Where to write the XLSX file")
    parser.add_argument(
        "--dpi",
        type=int,
        default=220,
        help="PDF rendering resolution used before table detection",
    )
    parser.add_argument(
        "--region-area-ratio",
        type=float,
        default=0.02,
        help="Minimum fraction of page area a detected table region must occupy",
    )
    parser.add_argument(
        "--line-coverage-ratio",
        type=float,
        default=0.2,
        help="How much of a row or column must be covered before it counts as a grid line",
    )
    parser.add_argument(
        "--min-cell-width",
        type=int,
        default=24,
        help="Ignore detected cells narrower than this many pixels",
    )
    parser.add_argument(
        "--min-cell-height",
        type=int,
        default=18,
        help="Ignore detected cells shorter than this many pixels",
    )
    parser.add_argument(
        "--cell-padding",
        type=int,
        default=3,
        help="Crop this many pixels away from each detected cell border before OCR",
    )
    parser.add_argument(
        "--ocr-score-threshold",
        type=float,
        default=0.35,
        help="Discard OCR snippets below this confidence score",
    )
    parser.add_argument(
        "--infer-numbers",
        action="store_true",
        help="Convert plain integer and decimal strings to numeric Excel cells",
    )
    parser.add_argument(
        "--log-level",
        default="INFO",
        choices=["DEBUG", "INFO", "WARNING", "ERROR"],
        help="Python logging verbosity",
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)

    logging.basicConfig(
        level=getattr(logging, args.log_level), format="%(levelname)s %(message)s"
    )

    if not args.input_pdf.exists():
        parser.error(f"Input PDF does not exist: {args.input_pdf}")

    tables = extract_tables(
        pdf_path=args.input_pdf,
        dpi=args.dpi,
        region_area_ratio=args.region_area_ratio,
        line_coverage_ratio=args.line_coverage_ratio,
        min_cell_width=args.min_cell_width,
        min_cell_height=args.min_cell_height,
        cell_padding=args.cell_padding,
        ocr_score_threshold=args.ocr_score_threshold,
    )
    write_workbook(
        tables=tables, output_path=args.output_xlsx, infer_numbers=args.infer_numbers
    )

    print(
        f"Extracted {len(tables)} table(s) from {args.input_pdf} into {args.output_xlsx}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
