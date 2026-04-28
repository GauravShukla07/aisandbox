#!/usr/bin/env python3
"""Extract table-like spreadsheet images from PDFs into Excel locally."""

from __future__ import annotations

import argparse
import csv
import re
import shutil
import sys
from dataclasses import dataclass
from itertools import pairwise
from pathlib import Path
from typing import Iterable, Sequence

import cv2
import numpy as np
import pypdfium2 as pdfium
import pytesseract
from openpyxl import Workbook
from PIL import Image

SUPPORTED_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}


@dataclass
class PageImage:
    name: str
    image: Image.Image


@dataclass
class ExtractionResult:
    name: str
    grid: list[list[str]]
    horizontal_lines: list[int]
    vertical_lines: list[int]


def parse_args(argv: Sequence[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Extract spreadsheet-like tables from scanned PDFs or images using "
            "OpenCV + Tesseract and write the result to XLSX/CSV."
        )
    )
    parser.add_argument("input_path", help="Path to a PDF, image, or directory of images.")
    parser.add_argument(
        "-o",
        "--output",
        help="Output XLSX path. Defaults to <input-stem>.xlsx in the current directory.",
    )
    parser.add_argument(
        "--csv-dir",
        help="Optional directory where one CSV file per page will be written.",
    )
    parser.add_argument(
        "--debug-dir",
        help="Optional directory for page overlays that show detected grid lines.",
    )
    parser.add_argument(
        "--dpi",
        type=int,
        default=300,
        help="PDF rendering DPI for scanned pages (default: 300).",
    )
    parser.add_argument(
        "--lang",
        default="eng",
        help="Tesseract language(s), e.g. eng or eng+deu (default: eng).",
    )
    parser.add_argument(
        "--psm",
        type=int,
        default=7,
        help="Tesseract page segmentation mode used per cell (default: 7).",
    )
    parser.add_argument(
        "--oem",
        type=int,
        default=3,
        help="Tesseract OCR engine mode used per cell (default: 3).",
    )
    return parser.parse_args(argv)


def ensure_tesseract() -> None:
    if shutil.which("tesseract") is None:
        raise SystemExit(
            "tesseract was not found on PATH. Install it locally first, for example:\n"
            "  sudo apt-get install tesseract-ocr tesseract-ocr-eng"
        )


def read_rgb_image(path: Path) -> Image.Image:
    with Image.open(path) as image:
        return image.convert("RGB")


def load_inputs(input_path: Path, dpi: int) -> list[PageImage]:
    if not input_path.exists():
        raise FileNotFoundError(f"Input path does not exist: {input_path}")

    if input_path.is_dir():
        images = [
            PageImage(name=path.stem, image=read_rgb_image(path))
            for path in sorted(input_path.iterdir())
            if path.suffix.lower() in SUPPORTED_IMAGE_SUFFIXES
        ]
        if not images:
            raise ValueError(f"No supported images found in directory: {input_path}")
        return images

    if input_path.suffix.lower() == ".pdf":
        return render_pdf(input_path, dpi=dpi)

    if input_path.suffix.lower() in SUPPORTED_IMAGE_SUFFIXES:
        return [PageImage(name=input_path.stem, image=read_rgb_image(input_path))]

    raise ValueError(
        "Unsupported input type. Pass a PDF, an image file, or a directory of images."
    )


def render_pdf(pdf_path: Path, dpi: int) -> list[PageImage]:
    scale = dpi / 72
    document = pdfium.PdfDocument(str(pdf_path))
    pages: list[PageImage] = []
    try:
        for index in range(len(document)):
            page = document[index]
            bitmap = page.render(scale=scale)
            pages.append(PageImage(name=f"page_{index + 1}", image=bitmap.to_pil().convert("RGB")))
    finally:
        document.close()
    return pages


def preprocess_image(image: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    blurred = cv2.GaussianBlur(gray, (3, 3), 0)
    binary = cv2.adaptiveThreshold(
        blurred,
        255,
        cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
        cv2.THRESH_BINARY_INV,
        31,
        15,
    )
    return gray, binary


def merge_positions(values: Iterable[float], tolerance: int) -> list[int]:
    sorted_values = sorted(values)
    if not sorted_values:
        return []

    clusters: list[list[float]] = [[sorted_values[0]]]
    for value in sorted_values[1:]:
        cluster = clusters[-1]
        center = sum(cluster) / len(cluster)
        if abs(value - center) <= tolerance:
            cluster.append(value)
        else:
            clusters.append([value])

    return [int(round(sum(cluster) / len(cluster))) for cluster in clusters]


def extract_line_positions(mask: np.ndarray, axis: str) -> list[int]:
    contours, _ = cv2.findContours(mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    positions: list[float] = []
    height, width = mask.shape[:2]

    for contour in contours:
        x, y, w, h = cv2.boundingRect(contour)
        if axis == "horizontal":
            if w < width * 0.25 or h > max(12, height * 0.1):
                continue
            positions.append(y + h / 2)
        else:
            if h < height * 0.25 or w > max(12, width * 0.1):
                continue
            positions.append(x + w / 2)

    if axis == "horizontal":
        tolerance = max(4, height // 200)
    else:
        tolerance = max(4, width // 200)
    return merge_positions(positions, tolerance=tolerance)


def detect_grid_lines(binary: np.ndarray) -> tuple[np.ndarray, np.ndarray, list[int], list[int]]:
    height, width = binary.shape[:2]
    horizontal_kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT, (max(16, width // 35), 1)
    )
    vertical_kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT, (1, max(16, height // 35))
    )

    horizontal = cv2.morphologyEx(binary, cv2.MORPH_OPEN, horizontal_kernel)
    vertical = cv2.morphologyEx(binary, cv2.MORPH_OPEN, vertical_kernel)
    horizontal = cv2.dilate(horizontal, cv2.getStructuringElement(cv2.MORPH_RECT, (3, 1)))
    vertical = cv2.dilate(vertical, cv2.getStructuringElement(cv2.MORPH_RECT, (1, 3)))

    horizontal_lines = extract_line_positions(horizontal, axis="horizontal")
    vertical_lines = extract_line_positions(vertical, axis="vertical")
    return horizontal, vertical, horizontal_lines, vertical_lines


def cleanup_text(text: str) -> str:
    compact_lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines()]
    compact_lines = [line for line in compact_lines if line]
    return "\n".join(compact_lines)


def build_grid_rectangles(horizontal_lines: Sequence[int], vertical_lines: Sequence[int]) -> list[list[tuple[int, int, int, int]]]:
    rectangles: list[list[tuple[int, int, int, int]]] = []
    for top, bottom in pairwise(horizontal_lines):
        row: list[tuple[int, int, int, int]] = []
        for left, right in pairwise(vertical_lines):
            row.append((left, top, right, bottom))
        rectangles.append(row)
    return rectangles


def ocr_cell(
    image: np.ndarray,
    rectangle: tuple[int, int, int, int],
    *,
    lang: str,
    psm: int,
    oem: int,
) -> str:
    left, top, right, bottom = rectangle
    inner_padding = max(2, min(right - left, bottom - top) // 18)
    crop = image[
        max(0, top + inner_padding) : max(0, bottom - inner_padding),
        max(0, left + inner_padding) : max(0, right - inner_padding),
    ]

    if crop.size == 0:
        return ""

    gray = cv2.cvtColor(crop, cv2.COLOR_BGR2GRAY)
    bordered = cv2.copyMakeBorder(
        gray, 10, 10, 10, 10, cv2.BORDER_CONSTANT, value=255
    )
    # Upscale small cell crops before OCR; Tesseract is much more reliable when
    # scanned spreadsheet text is closer to normal reading size.
    enlarged = cv2.resize(
        bordered, None, fx=2.0, fy=2.0, interpolation=cv2.INTER_CUBIC
    )
    normalized = cv2.threshold(
        enlarged, 0, 255, cv2.THRESH_BINARY | cv2.THRESH_OTSU
    )[1]
    config = f"--psm {psm} --oem {oem} -c preserve_interword_spaces=1"
    text = pytesseract.image_to_string(normalized, lang=lang, config=config)
    return cleanup_text(text)


def trim_empty_edges(grid: list[list[str]]) -> list[list[str]]:
    if not grid:
        return []

    nonempty_rows = [any(cell.strip() for cell in row) for row in grid]
    if not any(nonempty_rows):
        return []

    first_row = next(index for index, used in enumerate(nonempty_rows) if used)
    last_row = len(nonempty_rows) - next(
        index for index, used in enumerate(reversed(nonempty_rows)) if used
    )
    trimmed_rows = grid[first_row:last_row]

    width = max(len(row) for row in trimmed_rows)
    nonempty_columns = []
    for column_index in range(width):
        nonempty_columns.append(
            any(
                column_index < len(row) and row[column_index].strip()
                for row in trimmed_rows
            )
        )

    first_col = next(index for index, used in enumerate(nonempty_columns) if used)
    last_col = len(nonempty_columns) - next(
        index for index, used in enumerate(reversed(nonempty_columns)) if used
    )

    return [row[first_col:last_col] for row in trimmed_rows]


def extract_page(
    page: PageImage,
    *,
    lang: str,
    psm: int,
    oem: int,
    debug_dir: Path | None,
) -> ExtractionResult:
    image = cv2.cvtColor(np.array(page.image), cv2.COLOR_RGB2BGR)
    _, binary = preprocess_image(image)
    horizontal_mask, vertical_mask, horizontal_lines, vertical_lines = detect_grid_lines(binary)

    if len(horizontal_lines) < 2 or len(vertical_lines) < 2:
        raise ValueError(
            f"Could not detect a cell grid on {page.name}. "
            "This extractor expects visible row/column lines."
        )

    rectangles = build_grid_rectangles(horizontal_lines, vertical_lines)
    grid: list[list[str]] = []
    for row in rectangles:
        grid.append(
            [
                ocr_cell(image, rect, lang=lang, psm=psm, oem=oem)
                for rect in row
            ]
        )

    trimmed_grid = trim_empty_edges(grid)
    if debug_dir is not None:
        write_debug_overlay(
            page_name=page.name,
            image=image,
            horizontal_mask=horizontal_mask,
            vertical_mask=vertical_mask,
            horizontal_lines=horizontal_lines,
            vertical_lines=vertical_lines,
            debug_dir=debug_dir,
        )
    return ExtractionResult(
        name=page.name,
        grid=trimmed_grid,
        horizontal_lines=horizontal_lines,
        vertical_lines=vertical_lines,
    )


def write_debug_overlay(
    *,
    page_name: str,
    image: np.ndarray,
    horizontal_mask: np.ndarray,
    vertical_mask: np.ndarray,
    horizontal_lines: Sequence[int],
    vertical_lines: Sequence[int],
    debug_dir: Path,
) -> None:
    debug_dir.mkdir(parents=True, exist_ok=True)
    overlay = image.copy()

    for y in horizontal_lines:
        cv2.line(overlay, (0, y), (overlay.shape[1] - 1, y), (0, 200, 0), 1)
    for x in vertical_lines:
        cv2.line(overlay, (x, 0), (x, overlay.shape[0] - 1), (0, 0, 200), 1)

    combined_mask = cv2.bitwise_or(horizontal_mask, vertical_mask)
    blended_mask = cv2.applyColorMap(combined_mask, cv2.COLORMAP_OCEAN)
    preview = cv2.addWeighted(overlay, 0.75, blended_mask, 0.25, 0)
    cv2.imwrite(str(debug_dir / f"{page_name}_grid.png"), preview)


def sanitize_sheet_name(name: str, used_names: set[str]) -> str:
    clean = re.sub(r"[\[\]\*\?/:\\]", "_", name).strip() or "Sheet"
    clean = clean[:31]
    candidate = clean
    suffix = 1
    while candidate in used_names:
        extra = f"_{suffix}"
        candidate = f"{clean[:31 - len(extra)]}{extra}"
        suffix += 1
    used_names.add(candidate)
    return candidate


def write_xlsx(results: Sequence[ExtractionResult], output_path: Path) -> None:
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    used_names: set[str] = set()

    for result in results:
        sheet = workbook.create_sheet(sanitize_sheet_name(result.name, used_names))
        for row_index, row in enumerate(result.grid, start=1):
            for col_index, value in enumerate(row, start=1):
                sheet.cell(row=row_index, column=col_index, value=value)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_csvs(results: Sequence[ExtractionResult], csv_dir: Path) -> None:
    csv_dir.mkdir(parents=True, exist_ok=True)
    for result in results:
        csv_path = csv_dir / f"{result.name}.csv"
        with csv_path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerows(result.grid)


def default_output_path(input_path: Path) -> Path:
    return Path.cwd() / f"{input_path.stem}.xlsx"


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    ensure_tesseract()

    input_path = Path(args.input_path).expanduser().resolve()
    output_path = Path(args.output).expanduser().resolve() if args.output else default_output_path(input_path)
    csv_dir = Path(args.csv_dir).expanduser().resolve() if args.csv_dir else None
    debug_dir = Path(args.debug_dir).expanduser().resolve() if args.debug_dir else None

    pages = load_inputs(input_path, dpi=args.dpi)
    results = [
        extract_page(page, lang=args.lang, psm=args.psm, oem=args.oem, debug_dir=debug_dir)
        for page in pages
    ]

    write_xlsx(results, output_path)
    if csv_dir is not None:
        write_csvs(results, csv_dir)

    print(f"Wrote {len(results)} sheet(s) to {output_path}")
    if csv_dir is not None:
        print(f"Wrote CSV exports to {csv_dir}")
    if debug_dir is not None:
        print(f"Wrote debug overlays to {debug_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
